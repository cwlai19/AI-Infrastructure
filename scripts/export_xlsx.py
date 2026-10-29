"""Export the whole tracker to one Excel workbook for reference: every statement joined to its
source (title, link, grade), entity, and any analyst review; plus the resolved debt ledger,
sources, entities, links, reviews and monitor events."""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from tracker.model import DATA_DIR, load, resolve  # noqa: E402
from tracker.report import LEDGER_COLUMNS, _label, _value  # noqa: E402

store = load()
wb = Workbook()
HEAD = PatternFill("solid", fgColor="1F4E79")


def sheet(title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


readme = wb.active
readme.title = "README"
for line in [
    "Nscale credit tracker: full data export (generated from the statement files in the GitHub repo)",
    "",
    "Statements: one row per fact. Each fact has a source document, location in that document, the date it was true",
    "(effective_date), the date we recorded it (observed_date), how it was obtained (method) and a grade (A/B/C).",
    "Grades: A = filing, registry, rating agency or a lender's/counterparty's own filing; B = company press or an aggregator",
    "(e.g. CCIR), or an A source seen only as an excerpt; C = news, analysts, our own estimates.",
    "review_verdict: 'rejected' = kept on file but not used (our misreading); 'explained' = known, understood disagreement.",
    "Debt_ledger: the value shown for each facility after applying the resolution rules (latest date, then strongest source).",
    "Events: the daily monitor's log of new SEC and Companies House filings naming Nscale (status new / baseline / ignored).",
]:
    readme.append([line])
readme.column_dimensions["A"].width = 130

sources = store.sources
grade = {s.id: s.grade for s in store.statements}
sheet("Statements",
      ["id", "entity", "entity_type", "property", "value", "value_label", "unit", "grade", "method", "effective_date",
       "observed_date", "source_id", "source_type", "source_title", "locator", "source_url", "note", "review_verdict",
       "review_reason", "derived_from"],
      [[s.id, _label(store, s.entity), store.entities[s.entity]["type"], s.property, s.value,
        _value(store, s) if s.value else "", s.unit, s.grade, s.method, s.effective_date, s.observed_date, s.source,
        sources[s.source]["type"], sources[s.source]["title"], s.locator, sources[s.source].get("url", ""), s.note,
        store.reviews.get(s.id, {}).get("verdict", ""), store.reviews.get(s.id, {}).get("reason", ""),
        ";".join(s.derived_from)] for s in store.statements],
      [9, 34, 12, 24, 22, 30, 8, 6, 12, 12, 12, 22, 16, 40, 30, 40, 50, 10, 40, 14])

ledger_rows = []
for e, ent in store.entities.items():
    if ent["type"] != "facility":
        continue
    row = [_label(store, e)]
    for prop, _ in LEDGER_COLUMNS:
        r = resolve(store, e, prop)
        if r.status == "not_checked":
            row.append("")
        elif r.status == "undisclosed":
            row.append("undisclosed")
        else:
            shown = r.current if r.status == "multi" else r.current[:1]
            row.append("; ".join(f"{_value(store, s)} [{s.grade}]" for s in shown) + (" (disputed)" if r.conflicts else ""))
    ledger_rows.append(row)
sheet("Debt_ledger", ["Facility"] + [h for _, h in LEDGER_COLUMNS], ledger_rows, [40] + [24] * len(LEDGER_COLUMNS))

sheet("Sources", ["id", "type", "access", "publisher", "title", "url", "published_date", "retrieved_date", "note"],
      [[s.get(k, "") for k in ["id", "type", "access", "publisher", "title", "url", "published_date", "retrieved_date", "note"]]
       for s in sources.values()], [22, 18, 13, 30, 50, 60, 12, 12, 40])
sheet("Entities", ["id", "type", "display_name", "working_label", "note"],
      [[e["id"], e["type"], _label(store, e["id"]), e["label"], e.get("note", "")] for e in store.entities.values()],
      [26, 13, 45, 45, 50])
sheet("Entity_links", ["entity_a", "relation", "entity_b", "judged_by", "judged_date", "note"],
      [[_label(store, l["entity_a"]), l["relation"], _label(store, l["entity_b"]), l["judged_by"], l["judged_date"], l["note"]]
       for l in store.links], [40, 10, 40, 28, 12, 80])
sheet("Reviews", ["statement_id", "verdict", "reason", "reviewer", "date"],
      [[r["statement_id"], r["verdict"], r["reason"], r["reviewer"], r["date"]] for r in store.reviews.values()],
      [12, 10, 100, 28, 12])
events = list(csv.DictReader(open(DATA_DIR / "events.csv", encoding="utf-8"))) if (DATA_DIR / "events.csv").exists() else []
sheet("Events", ["event_date", "status", "feed", "kind", "subject", "title", "url"],
      [[e["event_date"], e["status"], e["feed"], e["kind"], _label(store, e["subject"]) if e["subject"] in store.entities else e["subject"],
        e["title"], e["url"]] for e in sorted(events, key=lambda e: e["event_date"], reverse=True)],
      [12, 10, 7, 14, 40, 70, 70])
for ws in wb.worksheets[1:]:
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)
out = Path(sys.argv[1] if len(sys.argv) > 1 else "reports/Nscale_tracker_data.xlsx")
wb.save(out)
print(f"wrote {out}: {len(store.statements)} statements, {len(ledger_rows)} facilities, {len(events)} events")
